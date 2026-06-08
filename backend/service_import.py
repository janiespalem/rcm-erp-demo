"""
Import usług z historycznego arkusza Excel.

Źródło zostaje mapowane na istniejące ProductTemplate + PriceHistory, bez
tworzenia osobnego obiegu poza Biuro.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utils import DEFAULT_MARGIN_PCT


@dataclass(frozen=True)
class ServiceImportRecord:
    source_row: int
    order_number: str
    client: str
    template: dict[str, Any]
    price_history: dict[str, Any]


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def parse_money(value: Any) -> float | None:
    text = _clean_text(value).lower()
    if not text:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    match = re.search(r"[-+]?\d[\d\s]*(?:[,.]\d+)?", text)
    if not match:
        return None
    number = match.group(0).replace(" ", "").replace(",", ".")
    try:
        return float(number)
    except ValueError:
        return None


def parse_hours(value: Any) -> float | None:
    text = _clean_text(value).lower()
    if not text:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    total = 0.0
    hours_match = re.search(r"(\d+(?:[,.]\d+)?)\s*h", text)
    minutes_match = re.search(r"(\d+(?:[,.]\d+)?)\s*min", text)
    if hours_match:
        total += float(hours_match.group(1).replace(",", "."))
    if minutes_match:
        total += float(minutes_match.group(1).replace(",", ".")) / 60
    if total:
        return round(total, 2)

    number = parse_money(text)
    return number


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def infer_service_category(description: str) -> str:
    text = description.lower()
    if any(word in text for word in ("remont", "napraw", "tulej")):
        return "remont"
    if "zbrojen" in text:
        return "zbrojenie"
    if "spaw" in text:
        return "spawanie"
    if any(word in text for word in ("cię", "ciecie", "cięcie")):
        return "cięcie"
    return "usługa"


def build_service_import_records(path: str | Path) -> list[ServiceImportRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    records: list[ServiceImportRecord] = []

    for row_index, row in enumerate(worksheet.iter_rows(min_row=4, values_only=True), start=4):
        if not any(value is not None for value in row):
            continue

        order_number = _clean_text(row[1] if len(row) > 1 else "")
        client = _clean_text(row[2] if len(row) > 2 else "")
        accepted_at = _as_date(row[3] if len(row) > 3 else None)
        finished_at = _as_date(row[5] if len(row) > 5 else None)
        description = _clean_text(row[6] if len(row) > 6 else "")
        material = _clean_text(row[7] if len(row) > 7 else "")
        material_cost = parse_money(row[8] if len(row) > 8 else None)
        quote_price = parse_money(row[9] if len(row) > 9 else None)
        total_price = parse_money(row[10] if len(row) > 10 else None)
        constructor_hours = parse_hours(row[11] if len(row) > 11 else None)
        production_hours = parse_hours(row[12] if len(row) > 12 else None)
        status = _clean_text(row[13] if len(row) > 13 else "")

        if not description and not order_number:
            continue

        category = infer_service_category(description)
        base_price = total_price or quote_price
        title = description[:90] if description else f"Usługa {order_number}"

        operations = []
        if constructor_hours:
            operations.append({"op": "Konstruktor", "wydział": "Konstruktor", "hours": constructor_hours, "rate_per_hour": 120})
        if production_hours:
            operations.append({"op": "Produkcja", "wydział": "Produkcja", "hours": production_hours, "rate_per_hour": 90})

        materials = []
        if material or material_cost:
            materials.append({
                "mat": material or "Materiał wg opisu",
                "qty": 1,
                "unit": "kpl",
                "cost_pln": material_cost,
            })

        notes_parts = [
            f"Źródło: Lista zleceń usługi, wiersz {row_index}",
            f"Numer zlecenia: {order_number}" if order_number else "",
            f"Klient: {client}" if client else "",
            f"Status: {status}" if status else "",
        ]

        template = {
            "name": title,
            "category": category,
            "operations_json": operations,
            "materials_json": materials,
            "instruction_blocks": [],
            "machines_json": [],
            "base_price_pln": base_price,
            "margin_pct": DEFAULT_MARGIN_PCT,
            "project_code": None,
            "position_nr": order_number or None,
            "notes": "\n".join(part for part in notes_parts if part),
        }
        price_history = {
            "order_type": category,
            "total_price_historical": base_price,
            "parameters_json": {
                "description": description,
                "material": material,
                "material_cost": material_cost,
                "constructor_hours": constructor_hours,
                "production_hours": production_hours,
                "finished_at": finished_at.isoformat() if finished_at else None,
                "source_row": row_index,
                "source_order_number": order_number,
            },
            "source": Path(path).name,
            "order_date": accepted_at or finished_at,
            "client": client,
        }
        records.append(ServiceImportRecord(row_index, order_number, client, template, price_history))

    return records
