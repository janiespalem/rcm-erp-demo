"""
Analytics, production queue, rentowność, harmonogram, benchmarks, service history, xlsx export.
"""
import io
import json
import logging
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from auth import require_role
from database import get_db
from models import Order, Quote, Setting, PriceHistory, OrderStatus

REVENUE_HISTORY_DAYS = 183
from schemas import (
    AnalyticsSummary, RevenueMonth, TopClient, OverdueOrder,
    BenchmarkOut, BenchmarkSample,
)
from utils import DEFAULT_LABOR_RATE_PLN

router = APIRouter()

_ALL     = require_role("biuro", "technolog", "ceo", "dyrektor_produkcji")
_MGMT    = require_role("ceo", "dyrektor_produkcji")
_TECH_UP = require_role("technolog", "ceo", "dyrektor_produkcji")

RENTOWNOSC_LIMIT    = 50
BENCHMARK_MIN_SAMPLES = 3
SERVICE_HISTORY_MAX = 500


@router.get("/api/production")
def get_production_queue(db: Session = Depends(get_db), _: dict = _TECH_UP):
    active_statuses = [OrderStatus.in_production, OrderStatus.w_trakcie, OrderStatus.gotowe]
    orders_with_quotes = (
        db.query(Order, Quote)
        .outerjoin(Quote, Quote.order_id == Order.id)
        .filter(Order.status.in_(active_statuses))
        .order_by(Order.deadline.asc().nullslast(), Order.id.desc())
        .all()
    )
    result = []
    for o, quote in orders_with_quotes:
        routing = []
        if quote and quote.processes_json:
            routing = [p.get("wydział") or p.get("name") or p.get("op", "")
                       for p in quote.processes_json if p]
        result.append({
            "id":           o.id,
            "order_number": o.order_number,
            "client":       o.client,
            "status":       o.status,
            "deadline":     o.deadline.strftime("%d.%m.%Y") if o.deadline else None,
            "description":  o.description,
            "material":     o.material,
            "routing":      [r for r in routing if r],
            "total_net":    float(quote.total_net) if quote and quote.total_net else None,
        })
    return result


@router.get("/api/rentownosc")
def get_rentownosc(db: Session = Depends(get_db), _: dict = _MGMT):
    labor_rate_row = db.query(Setting).filter(Setting.key == "labor_rate_pln").first()
    labor_rate = float(labor_rate_row.value) if labor_rate_row else DEFAULT_LABOR_RATE_PLN

    orders = (
        db.query(Order)
        .filter(Order.status.in_([OrderStatus.in_production, OrderStatus.w_trakcie, OrderStatus.gotowe, OrderStatus.wydane]))
        .order_by(Order.created_at.desc())
        .limit(RENTOWNOSC_LIMIT)
        .all()
    )
    result = []
    for o in orders:
        quote = db.query(Quote).filter(Quote.order_id == o.id).first()
        if not quote:
            continue
        cena = float(quote.total_net or 0)
        material_cost = float(quote.material_weight_kg or 0) * float(quote.material_price_per_kg or 0)
        if material_cost == 0:
            material_cost = float(quote.material_cost or 0)

        actual_hours_total = sum(
            float(op.actual_hours) for op in o.operations if op.actual_hours is not None
        )
        if actual_hours_total > 0:
            labor_cost = actual_hours_total * labor_rate
        else:
            labor_cost = float(quote.labor_hours or 0) * labor_rate
            if quote.processes_json:
                proc_hours = sum(float(p.get("hours") or 0) for p in quote.processes_json)
                labor_cost += proc_hours * labor_rate

        koszt = material_cost + labor_cost
        marza_pln = cena - koszt
        marza_pct = round((marza_pln / cena * 100), 1) if cena > 0 else None
        result.append({
            "id":           o.id,
            "order_number": o.order_number or f"ZW-{o.id}",
            "client":       o.client,
            "status":       o.status,
            "cena_pln":     round(cena, 2),
            "koszt_pln":    round(koszt, 2),
            "marza_pln":    round(marza_pln, 2),
            "marza_pct":    marza_pct,
            "actual_hours": actual_hours_total or None,
        })
    return result


@router.get("/api/analytics", response_model=AnalyticsSummary)
def get_analytics(db: Session = Depends(get_db), _: dict = _MGMT):
    total        = db.query(func.count(Order.id)).scalar()
    odrzut_count = db.query(func.count(Order.id)).filter(Order.triage_branch == "odrzut").scalar()
    standard     = db.query(func.count(Order.id)).filter(Order.triage_branch == "standard").scalar()
    niestandard  = db.query(func.count(Order.id)).filter(Order.triage_branch == "niestandard").scalar()
    in_prod      = db.query(func.count(Order.id)).filter(
        Order.status.in_([OrderStatus.in_production, OrderStatus.w_trakcie, OrderStatus.gotowe])
    ).scalar()
    done         = db.query(func.count(Order.id)).filter(Order.status.in_([OrderStatus.done, OrderStatus.wydane])).scalar()
    avg_margin   = db.query(func.avg(Quote.margin_pct)).scalar()

    quote_by_order = {q.order_id: q for q in db.query(Quote).all()}
    production_statuses = [
        OrderStatus.w_trakcie, OrderStatus.gotowe, OrderStatus.wydane,
        OrderStatus.done, OrderStatus.in_production,
    ]
    cutoff = datetime.combine(date.today() - timedelta(days=REVENUE_HISTORY_DAYS), datetime.min.time())
    monthly = defaultdict(lambda: {"orders": 0, "revenue": 0.0})
    production_orders = db.query(Order).filter(Order.status.in_(production_statuses)).all()
    for order in production_orders:
        if order.created_at and order.created_at < cutoff:
            continue
        month = order.created_at.strftime("%Y-%m") if order.created_at else date.today().strftime("%Y-%m")
        quote = quote_by_order.get(order.id)
        monthly[month]["orders"] += 1
        monthly[month]["revenue"] += float(quote.total_net or 0) if quote else 0.0
    revenue_by_month = [
        RevenueMonth(month=month, orders=data["orders"], revenue_pln=data["revenue"])
        for month, data in sorted(monthly.items())
    ]

    clients = defaultdict(lambda: {"orders": 0, "revenue": 0.0})
    active_orders = db.query(Order).filter(Order.status != OrderStatus.rejected).all()
    for order in active_orders:
        quote = quote_by_order.get(order.id)
        clients[order.client]["orders"] += 1
        clients[order.client]["revenue"] += float(quote.total_net or 0) if quote else 0.0
    top_clients = [
        TopClient(client=client, orders=data["orders"], revenue_pln=data["revenue"])
        for client, data in sorted(clients.items(), key=lambda x: (x[1]["orders"], x[1]["revenue"]), reverse=True)[:5]
    ]

    today = date.today()
    overdue = (
        db.query(Order)
        .filter(Order.deadline < today, Order.status.notin_([OrderStatus.done, OrderStatus.wydane, OrderStatus.rejected]))
        .order_by(Order.deadline.asc())
        .all()
    )
    overdue_orders = [
        OverdueOrder(
            id=o.id,
            order_number=o.order_number or "",
            client=o.client,
            status=o.status.value,
            deadline=o.deadline.isoformat() if o.deadline else "",
        )
        for o in overdue
    ]

    delivered_orders = (
        db.query(Order).filter(Order.delivered_at.isnot(None), Order.created_at.isnot(None)).all()
    )
    cycle_days_list = [
        (o.delivered_at - o.created_at).total_seconds() / 86400
        for o in delivered_orders if o.delivered_at and o.created_at
    ]
    avg_cycle_days = round(sum(cycle_days_list) / len(cycle_days_list), 1) if cycle_days_list else None

    q_to_start_list = [
        (o.started_at - o.quoted_at).total_seconds() / 86400
        for o in delivered_orders if o.started_at and o.quoted_at
    ]
    avg_quote_to_start = round(sum(q_to_start_list) / len(q_to_start_list), 1) if q_to_start_list else None

    accuracy_samples = []
    for o in delivered_orders:
        quote = quote_by_order.get(o.id)
        if not quote:
            continue
        planned = float(quote.labor_hours or 0)
        if quote.processes_json:
            planned += sum(float(p.get("hours") or 0) for p in quote.processes_json)
        actual = sum(float(op.actual_hours) for op in o.operations if op.actual_hours is not None)
        if planned > 0 and actual > 0:
            accuracy_samples.append(actual / planned * 100)
    estimate_accuracy = round(sum(accuracy_samples) / len(accuracy_samples), 1) if accuracy_samples else None

    return AnalyticsSummary(
        total_orders            = total or 0,
        odrzut_count            = odrzut_count or 0,
        odrzut_pct              = round((odrzut_count / total * 100) if total else 0, 1),
        standard_count          = standard or 0,
        niestandard_count       = niestandard or 0,
        avg_margin_pct          = round(float(avg_margin) * 100, 1) if avg_margin else None,
        orders_in_production    = in_prod or 0,
        orders_done             = done or 0,
        avg_cycle_days          = avg_cycle_days,
        avg_quote_to_start_days = avg_quote_to_start,
        estimate_accuracy_pct   = estimate_accuracy,
        revenue_by_month        = revenue_by_month,
        top_clients             = top_clients,
        overdue_orders          = overdue_orders,
    )


@router.get("/api/export/xlsx")
def export_xlsx(db: Session = Depends(get_db), _: dict = _MGMT):
    orders = db.query(Order).order_by(Order.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Zlecenia Demo"

    headers     = ["Nr", "Klient", "Status", "Gałąź", "Termin", "Wartość netto (PLN)", "Utworzono"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A3A5C")
    for col, h in enumerate(headers, 1):
        cell           = ws.cell(row=1, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, o in enumerate(orders, 2):
        quote = db.query(Quote).filter(Quote.order_id == o.id).first()
        ws.cell(row=row_idx, column=1, value=o.order_number)
        ws.cell(row=row_idx, column=2, value=o.client)
        ws.cell(row=row_idx, column=3, value=o.status.value if o.status else "")
        ws.cell(row=row_idx, column=4, value=o.triage_branch or "")
        ws.cell(row=row_idx, column=5, value=o.deadline.isoformat() if o.deadline else "")
        ws.cell(row=row_idx, column=6, value=float(quote.total_net) if quote and quote.total_net else 0)
        ws.cell(row=row_idx, column=7, value=o.created_at.strftime("%Y-%m-%d") if o.created_at else "")

    for col, width in zip(range(1, 8), [12, 25, 15, 12, 12, 20, 12]):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=zlecenia_rcm.xlsx"},
    )


@router.get("/api/harmonogram")
def get_harmonogram(db: Session = Depends(get_db), _: dict = _TECH_UP):
    orders = (
        db.query(Order)
        .filter(Order.status.notin_(["rejected", "done"]))
        .order_by(Order.deadline.asc())
        .all()
    )
    return [
        {
            "id":           o.id,
            "order_number": o.order_number,
            "client":       o.client,
            "status":       o.status,
            "deadline":     o.deadline.isoformat() if o.deadline else None,
            "branch":       o.triage_branch,
        }
        for o in orders
    ]


@router.get("/api/benchmarks/price-per-kg", response_model=BenchmarkOut)
def get_price_per_kg_benchmark(
    material: str,
    order_type: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = _TECH_UP,
):
    records = db.query(PriceHistory).all()
    samples = []
    for rec in records:
        params = rec.parameters_json or {}
        if isinstance(params, str):
            params = json.loads(params)
        record_material = str(params.get("material") or "").lower()
        if material and material.lower() not in record_material:
            continue
        if order_type and order_type.lower() not in str(rec.order_type or "").lower():
            continue
        weight = params.get("weight_kg", 0)
        pln_kg = params.get("pln_kg", 0)
        if weight and pln_kg:
            samples.append(BenchmarkSample(
                order_id=0,
                date=rec.order_date or date.today(),
                weight_kg=weight,
                total_net=float(rec.total_price_historical or 0),
                pln_kg=pln_kg,
            ))

    warning = None
    if len(samples) < BENCHMARK_MIN_SAMPLES:
        warning = f"Potrzeba minimum {BENCHMARK_MIN_SAMPLES} próbek dla wiarygodnego benchmarku"

    if not samples:
        return BenchmarkOut(avg_pln_kg=0.0, min_pln_kg=0.0, max_pln_kg=0.0, count=0,
                            warning=warning or "Brak danych", samples=[])

    pln_kg_values = [s.pln_kg for s in samples]
    return BenchmarkOut(
        avg_pln_kg=sum(pln_kg_values) / len(pln_kg_values),
        min_pln_kg=min(pln_kg_values),
        max_pln_kg=max(pln_kg_values),
        count=len(samples),
        warning=warning,
        samples=samples,
    )


@router.get("/api/service-history")
def list_service_history(limit: int = 100, db: Session = Depends(get_db), _: dict = _ALL):
    query = db.query(PriceHistory)
    preferred_source = "Kopia Lista zleceń usługi.xlsx"
    if query.filter(PriceHistory.source == preferred_source).first():
        query = query.filter(PriceHistory.source == preferred_source)

    rows = query.order_by(PriceHistory.id.desc()).limit(min(max(limit, 1), SERVICE_HISTORY_MAX)).all()
    out  = []
    for rec in rows:
        params = rec.parameters_json or {}
        if isinstance(params, str):
            try:
                params = json.loads(params)
            except json.JSONDecodeError:
                logging.getLogger(__name__).warning(
                    "Nieprawidłowy JSON w PriceHistory.parameters_json id=%s", rec.id
                )
                params = {}
        out.append({
            "id":                  rec.id,
            "order_date":          rec.order_date,
            "client":              rec.client,
            "order_type":          rec.order_type,
            "description":         params.get("description") or rec.order_type,
            "material":            params.get("material"),
            "material_cost":       params.get("material_cost"),
            "constructor_hours":   params.get("constructor_hours"),
            "production_hours":    params.get("production_hours"),
            "total_price":         float(rec.total_price_historical or 0),
            "source":              rec.source,
            "source_order_number": params.get("source_order_number"),
        })
    return out
